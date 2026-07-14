"""Text extraction for supported knowledge document types."""

import logging
from io import BytesIO
from pathlib import Path

from app.core.exceptions import ValidationError

logger = logging.getLogger(__name__)


class KnowledgeTextExtractor:
    def extract(self, filename: str, mime_type: str | None, data: bytes) -> str:
        suffix = Path(filename).suffix.lower()
        normalized_mime = (mime_type or "").lower()

        if suffix == ".txt" or normalized_mime.startswith("text/"):
            return self._decode_text(data)
        if suffix == ".pdf" or normalized_mime == "application/pdf":
            return self._extract_pdf(data)
        if suffix == ".docx" or normalized_mime.endswith("wordprocessingml.document"):
            return self._extract_docx(data)

        raise ValidationError("Only PDF, DOCX, and TXT knowledge files are supported")

    def _decode_text(self, data: bytes) -> str:
        for encoding in ("utf-8", "utf-16", "latin-1"):
            try:
                return data.decode(encoding).strip()
            except UnicodeDecodeError:
                continue
        raise ValidationError("Could not decode text file")

    def _extract_pdf(self, data: bytes) -> str:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ValidationError("PDF support requires the pypdf package") from exc

        reader = PdfReader(BytesIO(data))
        pages: list[str] = []
        failed_pages = 0

        for index, page in enumerate(reader.pages):
            try:
                pages.append(page.extract_text() or "")
            except Exception:
                # Some malformed pages raise instead of returning "".
                # Skip the page instead of failing the whole document.
                failed_pages += 1
                logger.warning("Failed to extract text from PDF page %d", index)

        if failed_pages:
            logger.warning("%d of %d PDF pages could not be read", failed_pages, len(reader.pages))

        text = "\n".join(pages).strip()
        if not text:
            raise ValidationError("No readable text was found in this PDF")
        return text

    def _extract_docx(self, data: bytes) -> str:
        try:
            from docx import Document
        except ImportError as exc:
            raise ValidationError("DOCX support requires the python-docx package") from exc

        document = Document(BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]

        # document.paragraphs does NOT include table content - without this,
        # any spec sheet / pricing table in a DOCX gets silently dropped.
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))

        text = "\n".join(part for part in parts if part.strip()).strip()
        if not text:
            raise ValidationError("No readable text was found in this DOCX file")
        return text